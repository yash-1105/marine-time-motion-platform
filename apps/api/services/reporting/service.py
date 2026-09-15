"""Phase 13 reporting orchestration built exclusively on ExecutiveDashboardService."""
from __future__ import annotations

import json
import os
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

from docx import Document
from pptx import Presentation
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from sqlalchemy import select
from sqlalchemy.orm import Session

from apps.api.core.config import settings
from apps.api.models.reporting import ReportArtifact, ReportDelivery, ReportRun, ReportSchedule, ReportTemplate
from apps.api.services.audit import log_audit_event
from apps.api.services.dashboard.executive import ExecutiveDashboardService

DAILY_SECTIONS = ["executive_summary", "movements", "anchorage", "pilotage", "towage", "berth", "cargo", "delays", "incidents", "resources", "exceptions", "attention_list", "methodology_provenance"]
TEMPLATES = {
    "daily-operations": ("Daily Operations", "IMPLEMENTED", DAILY_SECTIONS),
    "weekly-marine-performance": ("Weekly Marine Performance", "DEFERRED", ["turnaround", "pilot_tug", "berth", "delays", "resources", "bottlenecks", "targets", "history", "actions"]),
    "monthly-management-review": ("Monthly Management Review", "DEFERRED", ["kpis", "throughput", "turnaround", "delays", "utilisation", "criticality", "trends", "drivers", "risks", "recommendations"]),
    "quarterly-kpi-review": ("Quarterly KPI Review", "DEFERRED", ["scorecards", "targets", "trends", "variance", "rankings", "bottlenecks", "sustained_issues"]),
    "benchmark-performance": ("Benchmark Performance", "DEFERRED", ["internal", "historical", "peer", "gaps", "strengths", "opportunities", "recommendations"]),
}
CONTENT_TYPES = {"PDF": "application/pdf", "XLSX": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "PPTX": "application/vnd.openxmlformats-officedocument.presentationml.presentation", "DOCX": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}


def register_templates(db: Session) -> None:
    for template_id, (name, status, sections) in TEMPLATES.items():
        row = db.execute(select(ReportTemplate).where(ReportTemplate.template_id == template_id)).scalar_one_or_none()
        if not row:
            db.add(ReportTemplate(template_id=template_id, template_version="1.0", name=name, status=status, section_definitions=sections))
    db.commit()


class ReportService:
    def __init__(self, db: Session, tenant_id: str): self.db, self.tenant_id = db, tenant_id

    def build_result(self, run: ReportRun) -> dict[str, Any]:
        """Canonical data model. The sole numeric source is the governed dashboard service."""
        f = run.filters or {}
        governed = ExecutiveDashboardService(self.db, self.tenant_id).get_executive_summary(
            port_id=run.port_id, terminal_id=run.terminal_id, vessel_type=f.get("vessel_type"), cargo_type=f.get("cargo_type"),
            quality_status=f.get("quality_status"), start_date=run.period_start, end_date=run.period_end,
        )
        metrics = governed["lead_time_metrics"]
        unavailable = lambda reason: {"status": "UNAVAILABLE", "reason": reason}
        section = lambda title, value: {"title": title, "value": value}
        sections = {
            "executive_summary": section("Executive summary", governed["summary"]),
            "movements": section("Movements", {k: metrics[k] for k in ("Inward Movement", "Outward Movement")}),
            "anchorage": section("Anchorage", metrics["Anchorage Wait"]),
            "pilotage": section("Pilotage", {k: metrics[k] for k in ("Arrival Execution Delay", "Sailing Execution Delay")}),
            "towage": section("Towage", unavailable("No governed towage duration metric is available in the current dataset.")),
            "berth": section("Berth", metrics["Berth Stay"]),
            "cargo": section("Cargo", {"throughput": governed["throughput"], "working": metrics["Cargo Working"]}),
            "delays": section("Delays", governed["delays_summary"]),
            "incidents": section("Incidents", unavailable("No incident source records are available in the governed dataset.")),
            "resources": section("Resources", unavailable("Resource availability/roster inputs are not present in the governed dataset.")),
            "exceptions": section("Exceptions", {"quarantined_calls": governed["summary"]["quarantined_calls_count"], "outliers": governed["delays_summary"]["top_outliers"]}),
            "attention_list": section("Attention list", {"bottlenecks": governed["delays_summary"]["top_bottlenecks"], "outliers": governed["delays_summary"]["top_outliers"]}),
            "methodology_provenance": section("Methodology and provenance", {"source": "ExecutiveDashboardService", "lineage": governed["lineage"], "quarantine_disclosure": f"{governed['summary']['quarantined_calls_count']} quarantined calls are excluded from governed KPI populations by default."}),
        }
        return {"metadata": {"report_run_id": run.report_run_id, "template_id": run.template_id, "template_version": run.template_version, "reporting_period": {"start": str(run.period_start), "end": str(run.period_end)}, "filters": f, "scope": {"tenant_id": run.tenant_id, "port_id": run.port_id, "terminal_id": run.terminal_id}, "generated_at": datetime.now(UTC).isoformat(), "generated_by": run.requested_by or "system", "application_version": settings.application_version, "formula_versions": sorted({m.get("formula_version", "1.0") for m in metrics.values()}), "dataset_identity": governed["lineage"], "data_quality_status": governed["summary"], "synthetic_label": "SYNTHETIC / HISTORICAL DATA — NOT LIVE OPERATIONS" if governed["lineage"].get("dataset_type") == "HISTORICAL_SYNTHETIC" else None}, "sections": sections, "governed_dashboard": governed}

    def _lines(self, result: dict[str, Any]) -> list[str]:
        lines = ["Daily Operations Report", result["metadata"]["report_run_id"]]
        if result["metadata"]["synthetic_label"]: lines.append(result["metadata"]["synthetic_label"])
        lines.extend([f"Period: {result['metadata']['reporting_period']}", f"Filters: {result['metadata']['filters']}"])
        for key, sec in result["sections"].items(): lines.extend(["", sec["title"], json.dumps(sec["value"], default=str)[:1200]])
        return lines

    def _render(self, result: dict[str, Any], fmt: str, path: Path) -> None:
        lines = self._lines(result)
        if fmt == "PDF":
            c = canvas.Canvas(str(path), pagesize=letter); y = 750
            for line in lines:
                for fragment in [line[i:i+105] for i in range(0, len(line), 105)] or [""]:
                    c.drawString(36, y, fragment); y -= 13
                    if y < 36: c.showPage(); y = 750
            c.save(); return
        if fmt == "XLSX":
            from openpyxl import Workbook
            wb = Workbook(); ws = wb.active; ws.title = "Daily Operations"
            for i, line in enumerate(lines, 1): ws.cell(i, 1, line)
            wb.save(path); return
        if fmt == "DOCX":
            d = Document(); d.add_heading(lines[0], 0)
            for line in lines[1:]: d.add_paragraph(line)
            d.save(path); return
        p = Presentation(); slide = p.slides.add_slide(p.slide_layouts[1]); slide.shapes.title.text = lines[0]; slide.placeholders[1].text = "\n".join(lines[1:])[:7000]; p.save(path)

    def generate(self, run_id: str) -> ReportRun:
        run = self.db.execute(select(ReportRun).where(ReportRun.report_run_id == run_id, ReportRun.tenant_id == self.tenant_id)).scalar_one()
        if run.status == "CANCELLED": return run
        try:
            run.status, run.progress = "RUNNING", 10; self.db.commit()
            result = self.build_result(run); run.result_data, run.progress = result, 45; self.db.commit()
            base = Path(settings.report_storage_path) / run.tenant_id / run.report_run_id; base.mkdir(parents=True, exist_ok=True)
            for fmt in run.formats:
                path = base / f"daily-operations.{fmt.lower()}"; self._render(result, fmt, path)
                self.db.add(ReportArtifact(report_run_id=run.id, format=fmt, storage_key=str(path), content_type=CONTENT_TYPES[fmt], access_scope={"tenant_id": run.tenant_id, "port_id": run.port_id, "terminal_id": run.terminal_id}))
            run.status, run.progress = "AWAITING_APPROVAL", 100; self.db.commit()
        except Exception as exc:
            run.status, run.failure_reason = "FAILED", str(exc); self.db.commit(); raise
        return run

    def deliver(self, run: ReportRun, channel: str, recipient: str, fail: bool = False) -> ReportDelivery:
        delivery = ReportDelivery(report_run_id=run.id, channel=channel, recipient=recipient, attempts=1)
        if fail: delivery.status, delivery.error_message = "FAILED", "Configured delivery adapter failure"
        else: delivery.status, delivery.delivered_at = "DELIVERED", datetime.now(UTC)
        self.db.add(delivery); self.db.commit(); return delivery

    def retry_delivery(self, delivery: ReportDelivery) -> ReportDelivery:
        delivery.attempts += 1; delivery.status, delivery.error_message, delivery.delivered_at = "DELIVERED", None, datetime.now(UTC); self.db.commit(); return delivery
